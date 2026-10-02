"""Build the V02 disclosure, its seven diagrams, DOCX, and PDF from one source.

Owner: safety-constrained-coc. Requires python-docx, Pillow, and XeLaTeX.
Build intermediates are kept in /tmp; the supplied HWP is never edited.
"""

from pathlib import Path
import math
import re
import shutil
import subprocess
import tempfile

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "COC_INVENTION_DISCLOSURE_V02.md"
STEM = "coc_invention_disclosure_v02"
FONT = Path.home() / ".local/share/fonts/NotoSansCJK-Regular.ttc"
PDF_FONTS = ROOT.parent / "manuscript/latex-kiee-review-2022/fonts"


class Diagram:
    def __init__(self):
        self.image = Image.new("RGB", (1600, 1700), "white")
        self.draw = ImageDraw.Draw(self.image)
        self.font = ImageFont.truetype(str(FONT), 35)
        self.small = ImageFont.truetype(str(FONT), 29)

    def text(self, x, y, text, small=False):
        font = self.small if small else self.font
        self.draw.multiline_text((x, y), text, font=font, fill="black", anchor="mm", align="center", spacing=14)

    def box(self, x, y, w, h, text, number=""):
        self.draw.rounded_rectangle((x, y, x+w, y+h), radius=12, outline="black", width=4)
        for line in text.splitlines():
            assert self.draw.textlength(line, font=self.font) < w-28, (line, w)
        self.text(x+w/2, y+h/2+7, text)
        if number:
            self.draw.text((x+w-14, y+10), number, font=self.small, fill="black", anchor="rt")

    def arrow(self, points, label=None, at=None):
        self.draw.line(points, fill="black", width=4, joint="curve")
        x, y = points[-1]
        px, py = points[-2]
        angle = math.atan2(y-py, x-px)
        head = [(x, y)] + [(x-22*math.cos(angle+a), y-22*math.sin(angle+a)) for a in (-0.45, 0.45)]
        self.draw.polygon(head, fill="black")
        if label:
            self.text(*at, label, small=True)

    def save(self, number):
        path = ROOT / "figures" / f"coc_disclosure_fig{number:02d}_v02.png"
        path.parent.mkdir(exist_ok=True)
        self.image.save(path, dpi=(300, 300))


def figures():
    d = Diagram()
    for x, title, number in [(70, "장면 데이터", "110"), (450, "기본 CoC\n원인·행동목표", "120"), (830, "실행 가드", "130"), (1210, "후보·측정값", "140")]:
        d.box(x, 90, 320, 170, title, number)
    d.box(450, 380, 700, 160, "CoC + 실행 가드\n입력 문맥 구성", "150")
    d.arrow([(610,260),(610,380)]); d.arrow([(990,260),(990,380)])
    d.box(70, 740, 650, 200, "목표·가드 공동 판정\n허용 후보의 효용 비교", "160–190")
    d.box(880, 740, 650, 200, "장면 + 결합 문맥 + 후보\n멀티모달 모델 학습", "210")
    d.arrow([(800,540),(800,640),(395,640),(395,740)], "목표·가드", (540,610))
    d.arrow([(1370,260),(1570,260),(1570,610),(650,610),(650,740)])
    d.arrow([(230,260),(30,260),(30,680),(1050,680),(1050,740)])
    d.arrow([(1050,540),(1050,580),(1210,580),(1210,740)])
    d.arrow([(1370,260),(1370,740)])
    d.box(70,1100,650,160,"정답 후보 식별자\n감독 신호 생성","200")
    d.arrow([(395,940),(395,1100)])
    d.arrow([(720,1180),(790,1180),(790,1020),(1205,1020),(1205,940)],"정답 응답",(985,1080))
    d.box(880,1320,650,160,"학습 완료 모델","220")
    d.arrow([(1450,940),(1450,1320)],"파라미터 갱신",(1310,1220))
    d.text(800,1600,"기본 합성 구현과 일반 판정 설계의 관계: 본문 5.4–5.5절",True)
    d.save(1)

    d = Diagram()
    d.box(100,110,580,190,"관찰된 원인", "121"); d.box(920,110,580,190,"달성할 행동목표", "122")
    d.box(320,480,960,200,"기본 CoC R\n원인과 행동목표의 연결", "120")
    d.arrow([(390,300),(390,390),(600,390),(600,480)]); d.arrow([(1210,300),(1210,390),(1000,390),(1000,480)])
    d.box(100,860,580,250,"기본 CoC\n보행자에게 양보한 뒤\n경로를 진행한다", "120")
    d.box(920,860,580,250,"실행 가드 G\n상충 구역이 연속 1.5초\n비어 있을 때 진입", "130")
    d.arrow([(800,680),(800,770),(390,770),(390,860)])
    d.box(250,1320,1100,210,"S-C CoC = CONCAT(R, G)\n목표와 실행 허용조건을 하나의 문맥으로 결합", "150")
    d.arrow([(390,1110),(390,1320)]);d.arrow([(1210,1110),(1210,1320)])
    d.save(2)

    d = Diagram()
    d.draw.rounded_rectangle((50,30,1550,850),radius=16,outline="black",width=2)
    items=[("활성화\n의무 시작","131"),("불변\n유지할 상태","132"),("수치경계\n거리·속도·시간","133"),("금지\n허용하지 않는 행동","134"),("해제\n대기 종료","135"),("대체\n불확실·충족 불가","136"),("종료\n행동목표 완료","137")]
    for i,(label,num) in enumerate(items):
        row,col=divmod(i,3); d.box(90+col*500,70+row*270,420,200,label,num)
    d.box(90,920,1420,240,"장면·목표에 관련된 가드 인스턴스 선택\n적용 대상 · 기준값 · 단위 · 비교 방향 · 적용 구간", "138")
    d.arrow([(800,850),(800,920)])
    d.box(250,1330,1100,190,"선택한 가드 인스턴스의 논리곱\n모든 필수 가드를 만족할 때 Sat = 1", "139")
    d.arrow([(800,1160),(800,1330)])
    d.text(800,1610,"기본 실시예: 판정 가능한 가드가 미리 주어짐",True); d.save(3)

    d = Diagram()
    d.box(390,60,820,150,"후보 집합 · 목표 · 가드 · 효용", "140")
    d.box(300,330,1000,180,"Goal = 1 AND Sat = 1\n목표와 가드를 공동 충족하는가?", "170")
    d.arrow([(800,210),(800,330)])
    d.box(70,660,620,190,"공동 충족 후보 집합 F", "180")
    d.box(970,660,560,190,"정답 선택 대상에서 제외")
    d.arrow([(560,510),(560,580),(380,580),(380,660)],"예",(320,550))
    d.arrow([(1080,510),(1080,580),(1250,580),(1250,660)],"아니오",(1320,550))
    d.box(70,1020,620,160,"F가 비어 있는가?")
    d.arrow([(380,850),(380,1020)])
    d.box(970,1020,560,190,"정답 생성 불가\n학습 표본 제외", "201")
    d.arrow([(690,1100),(970,1100)],"예",(820,1050))
    d.box(70,1380,620,190,"최대 효용 후보 선택\n동률은 물리 후보 순서 적용", "190")
    d.arrow([(380,1180),(380,1380)],"아니오",(480,1270))
    d.box(970,1380,560,190,"정답 식별자 기록", "200")
    d.arrow([(690,1475),(970,1475)])
    d.text(800,1640,"일반 후보 집합의 설계 절차 · Goal과 Sat의 검사 순서는 제한하지 않음",True); d.save(4)

    d = Diagram()
    d.box(120,70,1360,190,"장면 x + CONCAT(R, G) + 후보 설명·측정값\n프롬프트 위치는 손실 계산에서 마스킹", "110·150·140")
    d.box(120,450,730,180,"멀티모달 모델\n후보 식별자 응답 확률", "210·211")
    d.arrow([(485,260),(485,450)])
    d.box(1030,450,450,180,"정답 응답 y", "200")
    d.box(240,850,1120,210,"정답 응답 토큰의 지도학습 손실\nL = −Σ log p(정답 토큰 | 입력, 이전 토큰)", "212")
    d.arrow([(485,630),(485,730),(650,730),(650,850)])
    d.arrow([(1255,630),(1255,730),(1050,730),(1050,850)])
    d.box(240,1270,1120,180,"LoRA 파라미터 갱신", "213")
    d.arrow([(800,1060),(800,1270)])
    d.arrow([(240,1360),(60,1360),(60,540),(120,540)])
    d.text(800,1570,"두 계약의 표본도 각각 감독학습 · 별도 쌍대·대조 손실 없음",True);d.save(5)

    d=Diagram()
    d.box(100,60,1400,210,"고정: 장면 x · 기본 CoC R · 후보 집합 T(x)\n현재 간격 7.2 m · A 즉시 변경 / B 중단 후 재시도", "300")
    d.box(90,450,640,210,"제1 가드: 최소 간격 5.5 m\n허용된 최대 효용 정답: A", "310")
    d.box(870,450,640,210,"제2 가드: 최소 간격 9.0 m\n허용된 목표 완료 정답: B", "320")
    d.arrow([(500,270),(500,450)]);d.arrow([(1100,270),(1100,450)])
    d.text(800,780,"후보 B는 재시도 시점에 9.0 m 이상임을 후보 속성으로 제공",True)
    d.box(90,1000,640,240,"학습 장면의 두 계약\n각 계약의 정답 토큰 감독\n개별 표본의 지도학습", "330")
    d.box(870,1000,640,240,"평가 장면의 두 계약\n두 정답을 모두 맞혔는지 확인\n계약 쌍 정확도 계산", "340")
    for start in (410,1190):
        for end in (410,1190):d.arrow([(start,660),(start,900),(end,900),(end,1000)])
    d.text(800,1450,"학습·평가는 분리된 장면 집합에서 각각 구성\n가드만 달라져 정답이 전환됨\n출력이 다르다는 사실만으로 평가 성공을 인정하지 않음",True);d.save(6)

    d=Diagram()
    d.box(70,50,660,200,"새 장면 · 결합 문맥\n복수 후보·측정값", "410·420·430")
    d.box(930,50,600,200,"구조화 가드 · 측정값\n대응 상태 시점", "420·430")
    d.box(70,400,660,160,"학습 완료 모델 → 선택 후보", "220·440")
    d.arrow([(400,250),(400,400)])
    d.box(440,780,720,200,"선택적 외부 검증\n목표 완료 및 가드 충족 판정", "450")
    d.arrow([(500,560),(500,680),(680,680),(680,780)],"검증 사용",(630,630))
    d.arrow([(1230,250),(1230,680),(1020,680),(1020,780)])
    d.box(70,1210,610,200,"후보 결과 출력\n검증 여부 구분", "460")
    d.box(900,1210,630,280,"부적합: 남은 후보 재검사\n공동 충족 후보 재선택\n없음·판정 불가: 실패 상태\n대체동작 또는 재계획 요청", "470")
    d.arrow([(600,980),(600,1080),(375,1080),(375,1210)],"공동 충족",(330,1030))
    d.arrow([(1020,980),(1020,1080),(1215,1080),(1215,1210)],"부적합·판정 불가",(1270,1040))
    d.arrow([(230,560),(30,560),(30,1310),(70,1310)])
    d.text(175,850,"검증\n미사용",True)
    d.arrow([(1215,1490),(1215,1530),(375,1530),(375,1410)],"재선택 성공",(750,1495))
    d.text(800,1630,"후보는 한 번씩 검사 · 실행 연계는 설계 실시예",True);d.save(7)


def blocks(source=SOURCE):
    lines=source.read_text(encoding="utf-8").splitlines()
    i=0
    while i<len(lines):
        line=lines[i]
        if not line:
            i+=1;continue
        if line.startswith("```"):
            i+=1; chunk=[]
            while not lines[i].startswith("```"):
                chunk.append(lines[i]);i+=1
            yield "code", "\n".join(chunk)
        elif line.startswith("|"):
            rows=[]
            while i<len(lines) and lines[i].startswith("|"):
                cells=[c.strip() for c in lines[i].strip("|").split("|")]
                if not all(re.fullmatch(r"[-:]+",c) for c in cells):rows.append(cells)
                i+=1
            yield "table",rows
            continue
        elif line=="<!-- pagebreak -->":yield "break", ""
        elif line.startswith("!["):
            yield "image",re.search(r"\((.*?)\)",line).group(1)
        elif line.startswith("#"):
            mark, value=line.split(" ",1);yield "h"+str(len(mark)),value
        else:yield "p",line
        i+=1


def docx_document(content, stem=STEM, title="발명내용설명서", version="V02"):
    doc=Document()
    section=doc.sections[0]
    section.page_width=Cm(21);section.page_height=Cm(29.7)
    section.top_margin=Cm(2);section.bottom_margin=Cm(2)
    section.left_margin=Cm(2);section.right_margin=Cm(2)
    for name,size in [("Normal",10.5),("Title",24),("Heading 1",15),("Heading 2",12)]:
        style=doc.styles[name];style.font.name="Noto Sans CJK KR";style.font.size=Pt(size)
        style.font.color.rgb=RGBColor(0,0,0)
        style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"),"Noto Sans CJK KR")
        style.paragraph_format.space_after=Pt(7)
    doc.styles["Normal"].paragraph_format.line_spacing=1.25
    section.header.paragraphs[0].text=f"{title}  |  실행 가드 통합형 CoC  |  {version}"
    section.header.paragraphs[0].runs[0].font.size=Pt(8)
    foot=section.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.CENTER
    fld=OxmlElement("w:fldSimple");fld.set(qn("w:instr"),"PAGE");foot._p.append(fld)
    for kind,value in content:
        if kind=="break":doc.add_page_break()
        elif kind=="image":doc.add_picture(str(ROOT/value),width=Cm(16.7))
        elif kind.startswith("h"):
            p=doc.add_paragraph(value,"Title" if kind=="h1" else "Heading 1" if kind=="h2" else "Heading 2")
            p.paragraph_format.keep_with_next=True
        elif kind=="table":
            table=doc.add_table(rows=0,cols=len(value[0]));table.style="Table Grid"
            for idx,row in enumerate(value):
                cells=table.add_row().cells
                props=table.rows[-1]._tr.get_or_add_trPr();props.append(OxmlElement("w:cantSplit"))
                if idx==0:props.append(OxmlElement("w:tblHeader"))
                for cell,text in zip(cells,row):
                    cell.text=text
                    for para in cell.paragraphs:
                        para.paragraph_format.space_after=Pt(4)
                        para.paragraph_format.keep_with_next=idx<len(value)-1
                        for run in para.runs:run.font.size=Pt(8);run.bold=idx==0
            doc.add_paragraph()
        else:
            p=doc.add_paragraph(value)
            if kind=="code":
                p.paragraph_format.left_indent=Cm(.3)
                p.paragraph_format.keep_together=True
                for run in p.runs:run.font.size=Pt(9)
    doc.core_properties.title=f"{title} — 실행 가드 통합형 CoC"
    doc.core_properties.subject=f"기술 검토용 {version}"
    doc.core_properties.author=""
    doc.save(ROOT/(stem+".docx"))


def escape(text):
    mapping={"\\":r"\textbackslash{}","&":r"\&","%":r"\%","$":r"\$","#":r"\#","_":r"\_\allowbreak{}","{":r"\{","}":r"\}","~":r"\textasciitilde{}","^":r"\textasciicircum{}"}
    return "".join(mapping.get(c,c) for c in text)


def pdf_document(content, stem=STEM, title="발명내용설명서", version="V02"):
    preamble=r"""\documentclass[10pt,a4paper]{article}
\usepackage{fontspec,geometry,graphicx,longtable,array,fancyhdr}
\usepackage[hidelinks]{hyperref}
\geometry{left=20mm,right=20mm,top=22mm,bottom=20mm,headheight=14pt}
\setmainfont[Path=FONTPATH/,Extension=.otf,UprightFont=*-Regular,BoldFont=*-Bold]{NotoSerifCJKkr}
\XeTeXlinebreaklocale "ko"
\XeTeXlinebreakskip=0pt plus 1pt
\setlength{\parindent}{0pt}\setlength{\parskip}{6pt}
\renewcommand{\baselinestretch}{1.25}
\setlength{\emergencystretch}{3em}
\setlength{\tabcolsep}{3pt}\renewcommand{\arraystretch}{1.25}
\pagestyle{fancy}\fancyhf{}
\fancyhead[L]{\small 발명내용설명서 | 실행 가드 통합형 CoC | V02}
\fancyfoot[C]{\thepage}
\hypersetup{pdftitle={발명내용설명서 — 실행 가드 통합형 CoC},pdfsubject={기술 검토용 수정본 V02}}
\begin{document}
""".replace("FONTPATH",str(PDF_FONTS)).replace("발명내용설명서",escape(title)).replace("V02",escape(version))
    out=[preamble]
    for kind,value in content:
        if kind=="break":out.append(r"\clearpage")
        elif kind=="image":out.append(r"\begin{center}\includegraphics[width=\linewidth,height=195mm,keepaspectratio]{"+str(ROOT/value)+r"}\end{center}")
        elif kind.startswith("h"):
            if kind=="h1":
                out.append(r"{\fontsize{24}{29}\selectfont\bfseries "+escape(value)+r"\par}")
            else:
                command=r"\section*" if kind=="h2" else r"\subsection*"
                out.append(command+"{"+escape(value)+"}")
        elif kind=="table":
            n=len(value[0]); widths={2:[.27,.73],3:[.23,.35,.42],5:[.30,.19,.19,.16,.16],7:[.08,.26,.07,.13,.11,.11,.24]}[n]
            total=sum(widths); available=481.9-6*n-0.4*(n+1)
            cols="|"+"|".join(r">{\raggedright\arraybackslash}p{"+f"{available*w/total:.2f}pt"+"}" for w in widths)+"|"
            out.append(r"\begin{center}\begin{minipage}{\linewidth}\fontsize{8}{11}\selectfont\begin{tabular}{"+cols+r"}\hline")
            head=" & ".join(r"\textbf{"+escape(v)+"}" for v in value[0])+r" \\ \hline"
            out.append(head)
            for row in value[1:]:out.append(" & ".join(escape(v) for v in row)+r" \\ \hline")
            out.append(r"\end{tabular}\end{minipage}\end{center}")
        elif kind=="code":
            out.append(r"\begin{center}\begin{minipage}{.97\linewidth}\fontsize{9}{13}\selectfont\setlength{\parskip}{0pt} "+r"\par ".join(escape(v) for v in value.splitlines())+r"\par\end{minipage}\end{center}")
        else:
            parts=re.split(r"(https://\S+)",value)
            out.append("".join(r"\url{"+s+"}" if s.startswith("https://") else escape(s) for s in parts)+r"\par")
    out.append(r"\end{document}")
    build=Path(tempfile.mkdtemp(prefix=stem+"_"))
    tex=build/(stem+".tex");tex.write_text("\n".join(out),encoding="utf-8")
    for _ in range(2):
        completed=subprocess.run(["xelatex","-interaction=nonstopmode","-halt-on-error",tex.name],cwd=build,text=True,stdout=subprocess.PIPE,stderr=subprocess.STDOUT)
        (build/"build_output.txt").write_text(completed.stdout)
        if completed.returncode:
            raise RuntimeError(f"XeLaTeX failed: {build}/build_output.txt\n{completed.stdout[-3000:]}")
    shutil.copyfile(build/(stem+".pdf"),ROOT/(stem+".pdf"))
    print(f"PDF build log: {build}")


if __name__=="__main__":
    figures()
    content=list(blocks())
    docx_document(content)
    pdf_document(content)
    print("Created seven diagrams, editable DOCX, and PDF from", SOURCE.name)
